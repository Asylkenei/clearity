"""Локальный каталог учебных данных. Модуль не используется при проверке нового видео."""

import argparse
import hashlib
import json
import math
import re
from collections import Counter
from pathlib import Path

import cv2
from openpyxl import load_workbook


def normalize_id(name: str) -> str:
    match = re.fullmatch(r"[XХ]-(\d+)", Path(str(name).strip()).stem.upper())
    if not match:
        raise ValueError(f"Неизвестный формат имени: {name}")
    return f"X-{int(match.group(1))}"


def parse_timestamp(value) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        seconds = float(value)
    else:
        text = str(value).strip().lower()
        clock = re.fullmatch(r"(\d+):(\d{1,2})", text)
        words = re.fullmatch(r"(?:(\d+)\s*минут\w*\s*)?(\d+(?:[.,]\d+)?)\s*секунд\w*", text)
        if clock:
            if int(clock.group(2)) >= 60:
                raise ValueError(f"Некорректный таймкод: {value}")
            seconds = int(clock.group(1)) * 60 + int(clock.group(2))
        elif words:
            seconds = int(words.group(1) or 0) * 60 + float(words.group(2).replace(",", "."))
        else:
            raise ValueError(f"Не удалось разобрать таймкод: {value}")
    if not math.isfinite(seconds) or seconds < 0:
        raise ValueError(f"Некорректный таймкод: {value}")
    return seconds


def build_manifest(directory: Path, labels_path: Path) -> dict:
    """Читаем таблицу и параметры файлов. Ни один кадр не отправляется наружу."""
    files = {}
    for path in directory.glob("*.mov"):
        video_id = normalize_id(path.name)
        if video_id in files:
            raise ValueError(f"Два файла соответствуют одному ID {video_id}")
        files[video_id] = path

    book = load_workbook(labels_path, read_only=True, data_only=True)
    records, warnings, seen = [], [], set()
    try:
        for name, timestamp, label in book["labels"].iter_rows(min_row=2, max_col=3, values_only=True):
            if name is None:
                continue
            video_id = normalize_id(name)
            if video_id in seen:
                raise ValueError(f"Повторная метка {video_id}")
            seen.add(video_id)
            if video_id not in files:
                warnings.append(f"Нет файла для {name}")
                continue
            label_text = str(label).strip().lower()
            if label_text not in ("нарушение", "без нарушения"):
                raise ValueError(f"Неизвестная метка для {name}: {label}")
            path = files[video_id]
            capture = cv2.VideoCapture(str(path))
            try:
                fps = capture.get(cv2.CAP_PROP_FPS)
                count = capture.get(cv2.CAP_PROP_FRAME_COUNT)
                ok, frame = capture.read()
                duration = count / fps if fps > 0 else None
                height, width = frame.shape[:2] if ok else (None, None)
            finally:
                capture.release()
            seconds = parse_timestamp(timestamp)
            if not ok:
                warnings.append(f"Не читается первый кадр {path.name}")
            if seconds is not None and duration is not None and seconds >= duration:
                warnings.append(f"Таймкод за пределами расчётной длительности: {path.name}")
            if label_text == "нарушение" and seconds is None:
                warnings.append(f"Не указан таймкод нарушения: {path.name}")
            with path.open("rb") as source:
                digest = hashlib.file_digest(source, "sha256").hexdigest()
            records.append({
                "video_id": video_id, "filename": path.name, "path": str(path.resolve()),
                "label": "violation" if label_text == "нарушение" else "no_violation",
                "label_timecode_sec": seconds, "event_type": None,
                "size_bytes": path.stat().st_size, "duration_sec_estimate": round(duration, 3) if duration else None,
                "fps": round(fps, 3), "width": width, "height": height,
                "first_frame_readable": bool(ok), "sha256": digest,
            })
    finally:
        book.close()
    for video_id in files.keys() - seen:
        warnings.append(f"Нет разметки для {files[video_id].name}")
    hashes = Counter(record["sha256"] for record in records)
    duplicates = [digest for digest, count in hashes.items() if count > 1]
    duplicate_groups = []
    for digest in duplicates:
        group = [record for record in records if record["sha256"] == digest]
        different_times = len({r["label_timecode_sec"] for r in group}) > 1
        conflicting_labels = len({r["label"] for r in group}) > 1
        duplicate_groups.append({
            "sha256": digest, "files": [r["filename"] for r in group],
            "label_timecodes_sec": [r["label_timecode_sec"] for r in group],
            "different_timecodes": different_times, "conflicting_labels": conflicting_labels,
        })
        if different_times or conflicting_labels:
            warnings.append("У одинаковых файлов разные отметки; проверьте вручную: " + ", ".join(r["filename"] for r in group))
    counts = Counter(record["label"] for record in records)
    return {
        "summary": {"videos": len(records), "violation": counts["violation"],
                    "no_violation": counts["no_violation"], "duplicate_hash_groups": len(duplicates),
                    "first_frames_readable": sum(r["first_frame_readable"] for r in records)},
        "warnings": warnings, "duplicate_hashes": duplicates, "duplicate_groups": duplicate_groups,
        "records": sorted(records, key=lambda item: int(item["video_id"].split("-")[1])),
        "note": "Типы событий отсутствуют. Длительность приблизительная. Проверен первый кадр, а не каждый кадр. Разделение на разработку и проверку ещё не выполнено.",
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Локально проверить каталог датасета")
    parser.add_argument("directory", type=Path)
    parser.add_argument("--labels", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    manifest = build_manifest(args.directory, args.labels or args.directory / "open_labels.xlsx")
    output = args.output or args.directory / "manifest.json"
    output.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest["summary"], ensure_ascii=False))
    print("Warnings:", manifest["warnings"])
    print("Saved:", output.resolve())
